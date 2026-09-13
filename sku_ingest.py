"""
SKU明细 入库脚本
──────────────────────────────────────────────
读取 Excel 的「SKU明细(追溯用)」子表 → 灌进 history_data.db 的 sku_detail 表。
底表换了新 Excel，就重跑一次这个脚本即可。

用法：
    python3 sku_ingest.py "/path/to/大促历史到手价汇总.xlsx"
    python3 sku_ingest.py            # 用默认路径
"""

import sqlite3
import sys
from pathlib import Path

import openpyxl

_HERE = Path(__file__).parent
DB_PATH = _HERE / 'history_data.db'
DEFAULT_XLSX = '/Users/iivyli/Downloads/大促历史到手价汇总_11月-3月-6月-8月-3.xlsx'
SHEET = 'SKU明细(追溯用)'

import sku_matcher as M  # 复用期次映射


def ingest(xlsx_path: str, db_path: str = None, verbose=True):
    db_path = db_path or str(DB_PATH)
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f'找不到子表「{SHEET}」，现有子表：{wb.sheetnames}')
    ws = wb[SHEET]

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        std_name, pid, period, sku_raw, price, raw_val = (list(r) + [None] * 6)[:6]
        if not std_name:
            continue
        period_key = M.PERIOD_MAP.get(str(period).strip(), str(period).strip())
        try:
            price = float(price) if price not in (None, '') else None
        except (TypeError, ValueError):
            price = None
        rows.append((
            str(std_name).strip(),
            str(pid).strip() if pid is not None else '',
            period_key,
            str(sku_raw) if sku_raw is not None else '',
            price,
            str(raw_val) if raw_val is not None else '',
        ))

    conn = sqlite3.connect(db_path)
    conn.execute('DROP TABLE IF EXISTS sku_detail')
    conn.execute('''CREATE TABLE sku_detail(
        rowid INTEGER PRIMARY KEY AUTOINCREMENT,
        std_name TEXT,
        product_id TEXT,
        period TEXT,
        sku_raw TEXT,
        price REAL,
        raw_val TEXT
    )''')
    conn.executemany(
        'INSERT INTO sku_detail(std_name,product_id,period,sku_raw,price,raw_val) VALUES(?,?,?,?,?,?)',
        rows)
    conn.execute('CREATE INDEX idx_sku_pid ON sku_detail(product_id)')
    conn.execute('CREATE INDEX idx_sku_name ON sku_detail(std_name)')
    conn.commit()

    n_prod = conn.execute('SELECT COUNT(DISTINCT std_name) FROM sku_detail').fetchone()[0]
    n_id = conn.execute('SELECT COUNT(DISTINCT product_id) FROM sku_detail').fetchone()[0]
    conn.close()
    if verbose:
        print(f'✅ 入库完成 → {db_path}')
        print(f'   行数 {len(rows)} | 标准商品 {n_prod} | 商品ID {n_id}')
    return len(rows)


if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not Path(path).exists():
        raise SystemExit(f'Excel 不存在：{path}')
    ingest(path)
